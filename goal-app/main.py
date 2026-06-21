import os
import json
import re
import io
import zipfile
import smtplib
from email.message import EmailMessage
from datetime import datetime
from typing import List, Optional
from pathlib import Path

from fastapi import FastAPI, HTTPException, Request, Query
from fastapi.responses import HTMLResponse, JSONResponse, Response, PlainTextResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel
from groq import Groq
import weasyprint
from dotenv import load_dotenv
from docx import Document

import database

load_dotenv()

app = FastAPI(title="Goal-to-Execution System")

BASE_DIR = Path(__file__).parent
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "static")), name="static")
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))

GROQ_API_KEY = os.environ.get("GROQ_API_KEY", "YOUR_API_KEY_HERE")
groq_client = Groq(api_key=GROQ_API_KEY)

SYSTEM_PROMPT = """You are an elite execution coach for ambitious students.
Your job is not to motivate. Convert goals into execution.
Output these exact sections:
GOAL | CURRENT STATE ANALYSIS | GAP ANALYSIS | 
12-MONTH ROADMAP | THIS MONTH | THIS WEEK | TODAY | 
TOP 3 HIGHEST LEVERAGE ACTIONS | BIGGEST RISKS | 
SUCCESS METRICS | FIRST ACTION IN NEXT 30 MINUTES
Rules: Brutally practical. No generic advice. 
No motivational fluff. Limited time and resources. 
Optimize for results not activity."""


class GenerateRequest(BaseModel):
    goal: str
    current_situation: str


class ActionToggle(BaseModel):
    action_id: int
    checked: bool


class EmailReminder(BaseModel):
    report_id: int
    to_email: str


def parse_report_sections(text: str) -> dict:
    sections = {
        "GOAL": "",
        "CURRENT STATE ANALYSIS": "",
        "GAP ANALYSIS": "",
        "12-MONTH ROADMAP": "",
        "THIS MONTH": "",
        "THIS WEEK": "",
        "TODAY": "",
        "TOP 3 HIGHEST LEVERAGE ACTIONS": "",
        "BIGGEST RISKS": "",
        "SUCCESS METRICS": "",
        "FIRST ACTION IN NEXT 30 MINUTES": "",
    }
    section_names_escaped = [re.escape(s) for s in sections.keys()]
    section_names_escaped.append(r'ROADMAP')
    joined = '|'.join(section_names_escaped)
    header_re = re.compile(
        r'^(?:\d+[\.\)]\s*)?(?:\#+\s+)?(?:\*\*)?(' + joined + r')(?:\*\*)?(?:\s*[:|]\s*)?$',
        re.IGNORECASE
    )
    current_section = None
    lines = text.split("\n")
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        m = header_re.match(stripped)
        if m:
            matched_name = m.group(1).upper()
            if matched_name == "ROADMAP":
                matched_name = "12-MONTH ROADMAP"
            else:
                for k in sections:
                    if k.upper() == matched_name:
                        matched_name = k
                        break
            if matched_name in sections:
                current_section = matched_name
                content = stripped[m.end():].strip()
                if content and content not in ('|', ':'):
                    sections[current_section] = content
            continue
        if current_section:
            if sections[current_section]:
                sections[current_section] += "\n" + stripped
            else:
                sections[current_section] = stripped
    for k, v in sections.items():
        sections[k] = v.strip()
    return sections


def markdown_to_html_sections(sections: dict, report_id: int = 0) -> tuple:
    icon_map = {
        "GOAL": "🎯",
        "CURRENT STATE ANALYSIS": "📊",
        "GAP ANALYSIS": "🔍",
        "12-MONTH ROADMAP": "🗺️",
        "THIS MONTH": "📅",
        "THIS WEEK": "📋",
        "TODAY": "⚡",
        "TOP 3 HIGHEST LEVERAGE ACTIONS": "💎",
        "BIGGEST RISKS": "⚠️",
        "SUCCESS METRICS": "📈",
        "FIRST ACTION IN NEXT 30 MINUTES": "🔥",
    }
    priority_map = {
        "BIGGEST RISKS": "critical",
        "TOP 3 HIGHEST LEVERAGE ACTIONS": "high",
        "FIRST ACTION IN NEXT 30 MINUTES": "high",
        "GAP ANALYSIS": "high",
    }
    badge_map = {
        "BIGGEST RISKS": "CRITICAL",
        "TOP 3 HIGHEST LEVERAGE ACTIONS": "HIGH",
        "FIRST ACTION IN NEXT 30 MINUTES": "URGENT",
        "GAP ANALYSIS": "HIGH",
    }
    actionable_sections = ["THIS MONTH", "THIS WEEK", "TODAY", "TOP 3 HIGHEST LEVERAGE ACTIONS", "FIRST ACTION IN NEXT 30 MINUTES"]

    html_cards = []
    all_actions = []
    action_idx = 0

    for section_name, content in sections.items():
        if not content:
            continue
        icon = icon_map.get(section_name, "📌")
        badge = badge_map.get(section_name, "")
        priority = priority_map.get(section_name, "medium")
        is_urgent = section_name == "FIRST ACTION IN NEXT 30 MINUTES"
        formatted_content = content
        formatted_content = re.sub(r'\*\*(.*?)\*\*', r'<strong>\1</strong>', formatted_content)
        lines = formatted_content.split("\n")
        html_lines = []
        in_list = False
        is_actionable = section_name in actionable_sections
        for i, line in enumerate(lines):
            stripped = line.strip()
            if stripped.startswith(("- ", "* ", "• ")):
                item_text = stripped[2:].strip()
                if not in_list:
                    html_lines.append("<ul>")
                    in_list = True
                if is_actionable and report_id:
                    cb = f'<input type="checkbox" class="action-cb" data-idx="{action_idx}" data-report="{report_id}" data-section="{section_name}"> '
                    all_actions.append((report_id, section_name, item_text))
                    action_idx += 1
                    html_lines.append(f"<li>{cb}{item_text}</li>")
                else:
                    html_lines.append(f"<li>{item_text}</li>")
            else:
                if in_list:
                    html_lines.append("</ul>")
                    in_list = False
                if stripped:
                    html_lines.append(f"<p>{stripped}</p>")
                elif i < len(lines) - 1:
                    html_lines.append("<br/>")
        if in_list:
            html_lines.append("</ul>")
        content_html = "".join(html_lines)
        risk_meter_html = ""
        if section_name == "BIGGEST RISKS":
            risk_items = re.findall(r'(?:^|\n)\s*[-*•]\s*(.*?)(?=\n\s*[-*•]|\n\n|$)', content, re.DOTALL)
            if not risk_items:
                risk_items = [r.strip() for r in content.split("\n") if r.strip()][:3]
            risk_meter_html = '<div class="risk-meter-container">'
            for risk in risk_items[:4]:
                risk = risk.strip().lstrip("-*• ").strip()
                if risk:
                    risk_severity = "high" if len(risk) > 60 else "medium"
                    risk_width = 85 if risk_severity == "high" else 55
                    risk_meter_html += f"""
                    <div class="risk-item">
                        <div class="risk-label">{risk[:80]}{'...' if len(risk) > 80 else ''}</div>
                        <div class="risk-bar-bg">
                            <div class="risk-bar-fill severity-{risk_severity}" style="width: {risk_width}%"></div>
                        </div>
                    </div>"""
            risk_meter_html += "</div>"
        extra_class = "card-urgent" if is_urgent else ""
        html_cards.append({
            "title": section_name,
            "icon": icon,
            "content": content_html + risk_meter_html,
            "badge": badge,
            "priority": priority,
            "extra_class": extra_class,
            "is_urgent": is_urgent,
        })
    return html_cards, all_actions


def generate_markdown_report(goal: str, current_situation: str, sections: dict) -> str:
    md = f"""# Goal-to-Execution Report

**Goal:** {goal}

**Generated:** {datetime.now().strftime("%Y-%m-%d %H:%M")}

---
"""
    for section_name, content in sections.items():
        if content:
            md += f"\n## {section_name}\n\n{content}\n"
    return md


@app.on_event("startup")
def startup():
    database.init_db()


@app.get("/", response_class=HTMLResponse)
async def index(request: Request):
    return templates.TemplateResponse("index.html", {"request": request})


@app.post("/generate")
async def generate_report(req: GenerateRequest):
    try:
        completion = groq_client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            max_tokens=4096,
            temperature=0.7,
            messages=[
                {"role": "system", "content": SYSTEM_PROMPT},
                {
                    "role": "user",
                    "content": f"My goal is: {req.goal}\n\nMy current situation is: {req.current_situation}\n\nAnalyze my goal and current situation, then provide the exact execution sections requested."
                }
            ]
        )
        response_text = completion.choices[0].message.content
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Groq API error: {str(e)}")

    sections = parse_report_sections(response_text)
    cards, actions = markdown_to_html_sections(sections, report_id=0)
    markdown_report = generate_markdown_report(req.goal, req.current_situation, sections)

    report_json = {
        "goal": req.goal,
        "current_situation": req.current_situation,
        "sections": sections,
        "cards": cards,
        "markdown": markdown_report,
        "created_at": datetime.now().isoformat()
    }

    report_id = database.save_report(req.goal, req.current_situation, report_json, markdown_report)

    # Save actions
    for act in actions:
        database.save_actions(act[0], act[1], [act[2]])

    # Re-render cards with real report_id for checkboxes
    cards_with_cb, _ = markdown_to_html_sections(sections, report_id=report_id)

    return {
        "id": report_id,
        "goal": req.goal,
        "cards": cards_with_cb,
        "sections": sections,
        "created_at": report_json["created_at"]
    }


@app.get("/reports")
async def list_reports(search: Optional[str] = None):
    if search:
        reports = database.search_reports(search)
    else:
        reports = database.get_all_reports()
    return {"reports": reports}


@app.get("/reports/{report_id}")
async def get_report(report_id: int):
    report = database.get_report(report_id)
    if not report:
        raise HTTPException(status_code=404, detail="Report not found")
    # Re-render cards with report_id for checkboxes
    sections = report["report_json"]["sections"]
    cards_with_cb, _ = markdown_to_html_sections(sections, report_id=report_id)
    report_with_cards = {**report}
    report_with_cards["report_json"]["cards"] = cards_with_cb
    return report_with_cards


@app.delete("/reports/{report_id}")
async def delete_report(report_id: int):
    deleted = database.delete_report(report_id)
    if not deleted:
        raise HTTPException(status_code=404, detail="Report not found")
    return {"ok": True}


# ----- ACTION TRACKING -----

@app.post("/actions/toggle")
async def toggle_action(req: ActionToggle):
    result = database.toggle_action(req.action_id, req.checked)
    if not result:
        raise HTTPException(status_code=404, detail="Action not found")
    return result


@app.get("/actions/report/{report_id}")
async def get_actions(report_id: int):
    return {"actions": database.get_actions_for_report(report_id)}


@app.get("/actions/history")
async def action_history():
    return {"actions": database.get_action_history()}


# ----- EXPORTS -----

@app.get("/export/zip")
async def export_zip():
    reports = database.get_all_reports()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        for r in reports:
            report = database.get_report(r["id"])
            if report:
                md = report["report_markdown"]
                safe_name = f"report-{r['id']}-{r['goal'][:30].replace('/', '_').replace(' ', '_')}.md"
                zf.writestr(safe_name, md)
    buf.seek(0)
    return Response(
        content=buf.read(),
        media_type="application/zip",
        headers={"Content-Disposition": "attachment; filename=all-reports.zip"}
    )


@app.get("/export/{report_id}")
async def export_pdf(report_id: int):
    report = database.get_report(report_id)
    if not report:
        raise HTTPException(status_code=404, detail="Report not found")
    html_content = templates.TemplateResponse("report.html", {
        "request": None,
        "report": report,
        "sections": report["report_json"]["sections"],
    }).body.decode()
    pdf_bytes = weasyprint.HTML(string=html_content).write_pdf()
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=goal-report-{report_id}.pdf"}
    )


@app.get("/export/markdown/{report_id}")
async def export_markdown(report_id: int):
    report = database.get_report(report_id)
    if not report:
        raise HTTPException(status_code=404, detail="Report not found")
    md_content = report["report_markdown"]
    return PlainTextResponse(
        content=md_content,
        headers={"Content-Disposition": f"attachment; filename=goal-report-{report_id}.md"}
    )


@app.get("/export/docx/{report_id}")
async def export_docx(report_id: int):
    report = database.get_report(report_id)
    if not report:
        raise HTTPException(status_code=404, detail="Report not found")
    doc = Document()
    doc.add_heading("Goal-to-Execution Report", level=1)
    doc.add_paragraph(f"Goal: {report['goal']}")
    doc.add_paragraph(f"Generated: {report['created_at']}")
    doc.add_paragraph("---")
    sections = report["report_json"]["sections"]
    for section_name, content in sections.items():
        if content:
            doc.add_heading(section_name, level=2)
            for line in content.split("\n"):
                stripped = line.strip()
                if stripped.startswith(("- ", "* ", "• ")):
                    doc.add_paragraph(stripped[2:].strip(), style="List Bullet")
                elif stripped:
                    doc.add_paragraph(stripped)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename=goal-report-{report_id}.docx"}
    )


# ----- IMPORT -----

class ImportRequest(BaseModel):
    goal: str
    current_situation: str
    sections: dict


@app.post("/import")
async def import_report(req: ImportRequest):
    cards, actions = markdown_to_html_sections(req.sections, report_id=0)
    markdown_report = generate_markdown_report(req.goal, req.current_situation, req.sections)
    report_json = {
        "goal": req.goal,
        "current_situation": req.current_situation,
        "sections": req.sections,
        "cards": cards,
        "markdown": markdown_report,
        "created_at": datetime.now().isoformat()
    }
    report_id = database.save_report(req.goal, req.current_situation, report_json, markdown_report)
    for act in actions:
        database.save_actions(act[0], act[1], [act[2]])
    cards_with_cb, _ = markdown_to_html_sections(req.sections, report_id=report_id)
    return {
        "id": report_id,
        "goal": req.goal,
        "cards": cards_with_cb,
        "created_at": report_json["created_at"]
    }


@app.post("/import/json")
async def import_json(request: Request):
    body = await request.json()
    goal = body.get("goal", "Imported Goal")
    current_situation = body.get("current_situation", "")
    sections = body.get("sections", {})
    return await import_report(ImportRequest(goal=goal, current_situation=current_situation, sections=sections))


# ----- ANALYTICS -----

@app.get("/analytics")
async def analytics():
    return database.get_analytics()


# ----- CALENDAR -----

@app.get("/calendar")
async def calendar():
    return {"events": database.get_calendar_data()}


# ----- EMAIL REMINDERS -----

@app.post("/reminders/email")
async def send_email_reminder(req: EmailReminder):
    report = database.get_report(req.report_id)
    if not report:
        raise HTTPException(status_code=404, detail="Report not found")
    sections = report["report_json"]["sections"]
    today_actions = sections.get("TODAY", "Not specified")
    this_week = sections.get("THIS WEEK", "Not specified")
    smtp_server = os.environ.get("SMTP_SERVER", "smtp.gmail.com")
    smtp_port = int(os.environ.get("SMTP_PORT", "587"))
    smtp_user = os.environ.get("SMTP_USER", "")
    smtp_pass = os.environ.get("SMTP_PASS", "")
    if not smtp_user or not smtp_pass:
        raise HTTPException(status_code=400, detail="SMTP not configured. Set SMTP_USER and SMTP_PASS in .env")
    msg = EmailMessage()
    msg.set_content(f"""Goal-to-Execution - Daily Reminder

Goal: {report['goal']}

TODAY's Actions:
{today_actions}

THIS WEEK:
{this_week}

---
Generated by Goal-to-Execution System
""")
    msg["Subject"] = f"Goal Reminder: {report['goal'][:50]}"
    msg["From"] = smtp_user
    msg["To"] = req.to_email
    try:
        with smtplib.SMTP(smtp_server, smtp_port) as server:
            server.starttls()
            server.login(smtp_user, smtp_pass)
            server.send_message(msg)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Email failed: {str(e)}")
    return {"ok": True, "sent_to": req.to_email}


# ----- WEB PUSH / SERVICE WORKER -----

@app.get("/service-worker.js")
async def service_worker():
    sw_path = BASE_DIR / "static" / "service-worker.js"
    if sw_path.exists():
        return Response(content=sw_path.read_text(), media_type="application/javascript")
    return Response(content="", media_type="application/javascript")


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000)
